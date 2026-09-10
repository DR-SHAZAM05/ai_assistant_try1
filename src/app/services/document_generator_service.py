"""Document generation service for official UNITBV practice paperwork.

Generates:
1. Convenție-cadru privind efectuarea stagiului de practică (tripartite agreement)
2. Caiet de practică / Jurnal de activitate (weekly activity log & tutor evaluation)
in Microsoft Word (.docx) format with academic UNITBV visual standards.
"""

import io
from pathlib import Path
from typing import Any, Dict, List, Optional
import docx
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Inches, Pt, RGBColor

from src.app.core.config import settings
from src.app.core.logging import logger


class DocumentGeneratorService:
    """Service responsible for assembling and styling official UNITBV practice documents."""

    # Brand Colors
    COLOR_NAVY = RGBColor(0x1B, 0x36, 0x5D)      # #1B365D Academic Navy (UNITBV)
    COLOR_SLATE = RGBColor(0x2E, 0x5B, 0x88)     # #2E5B88 Secondary Slate Blue
    COLOR_TEAL = RGBColor(0x0D, 0x94, 0x88)      # #0D9488 Accent Teal
    COLOR_DARK = RGBColor(0x1E, 0x29, 0x3B)      # #1E293B Dark Charcoal
    COLOR_MUTED = RGBColor(0x64, 0x74, 0x8B)     # #64748B Muted Gray
    COLOR_WHITE = RGBColor(0xFF, 0xFF, 0xFF)

    HEX_NAVY = "1B365D"
    HEX_SLATE = "2E5B88"
    HEX_BORDER = "CBD5E1"
    HEX_LIGHT_BG = "F8FAFC"
    HEX_CALLOUT_BG = "F0F7FF"

    FONT_NAME = "Arial"

    def __init__(self):
        self._assets_dir = self._resolve_assets_dir()

    def _resolve_assets_dir(self) -> Path:
        """Find the path containing UNITBV / FIESC logo assets."""
        candidates = [
            Path("/app/docs/assets"),
            Path("docs/assets"),
            Path(__file__).resolve().parent.parent.parent.parent / "docs" / "assets",
        ]
        for c in candidates:
            if c.exists() and (c / "logo_unitbv.png").exists():
                return c
        return candidates[0]

    # -------------------------------------------------------------------------
    # XML Styling Helpers
    # -------------------------------------------------------------------------
    def _apply_font(
        self,
        run,
        size_pt: float = 10.5,
        color_rgb: Optional[RGBColor] = None,
        bold: bool = False,
        italic: bool = False,
    ):
        run.font.name = self.FONT_NAME
        run.font.size = Pt(size_pt)
        if color_rgb is not None:
            run.font.color.rgb = color_rgb
        run.bold = bold
        run.italic = italic

        rPr = run._r.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts")
            rPr.append(rFonts)
        rFonts.set(qn("w:ascii"), self.FONT_NAME)
        rFonts.set(qn("w:hAnsi"), self.FONT_NAME)
        rFonts.set(qn("w:cs"), self.FONT_NAME)

    def _set_cell_background(self, cell, fill_hex: str):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        tcPr.append(shd)

    def _set_cell_margins(self, cell, top=100, bottom=100, left=120, right=120):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(
            f'<w:tcMar {nsdecls("w")}>'
            f'<w:top w:w="{top}" w:type="dxa"/>'
            f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
            f'<w:left w:w="{left}" w:type="dxa"/>'
            f'<w:right w:w="{right}" w:type="dxa"/>'
            f"</w:tcMar>"
        )
        tcPr.append(tcMar)

    def _set_table_borders(self, table, color_hex: str = "CBD5E1"):
        tblPr = table._tbl.tblPr
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="6" w:space="0" w:color="{color_hex}"/>'
            f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{color_hex}"/>'
            f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color_hex}"/>'
            f'<w:insideV w:val="single" w:sz="4" w:space="0" w:color="{color_hex}"/>'
            f'<w:left w:val="single" w:sz="6" w:space="0" w:color="{color_hex}"/>'
            f'<w:right w:val="single" w:sz="6" w:space="0" w:color="{color_hex}"/>'
            f"</w:tblBorders>"
        )
        tblPr.append(borders)

    def _set_invisible_borders(self, table):
        tblPr = table._tbl.tblPr
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
            f'<w:bottom w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
            f'<w:insideH w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
            f'<w:insideV w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
            f'<w:left w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
            f'<w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
            f"</w:tblBorders>"
        )
        tblPr.append(borders)

    def _set_row_cant_split(self, row):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

    def _configure_document_page(self, doc: docx.Document, doc_title: str):
        """Setup A4 page margins (2.0 cm) and dynamic running headers and footers."""
        for section in doc.sections:
            section.page_width = Cm(21.0)
            section.page_height = Cm(29.7)
            section.top_margin = Cm(2.0)
            section.bottom_margin = Cm(2.0)
            section.left_margin = Cm(2.0)
            section.right_margin = Cm(2.0)
            section.header_distance = Cm(1.0)
            section.footer_distance = Cm(1.0)

            # Header
            hdr = section.header
            p_hdr = hdr.paragraphs[0]
            p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p_hdr.paragraph_format.space_after = Pt(2)
            r_hdr = p_hdr.add_run(f"UNITBV  |  Facultatea IESC  |  {doc_title}")
            self._apply_font(r_hdr, size_pt=8.0, color_rgb=self.COLOR_MUTED)

            # Footer
            ftr = section.footer
            p_ftr = ftr.paragraphs[0]
            p_ftr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p_ftr.paragraph_format.space_before = Pt(4)
            r_ftr_l = p_ftr.add_run("Universitatea Transilvania din Brașov • Document Oficial de Practică")
            self._apply_font(r_ftr_l, size_pt=8.0, color_rgb=self.COLOR_MUTED)

    def _add_institutional_header(self, doc: docx.Document):
        """Adds UNITBV and IESC branding with optional logos."""
        tbl = doc.add_table(rows=1, cols=3)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        self._set_invisible_borders(tbl)
        self._set_row_cant_split(tbl.rows[0])

        c_left, c_center, c_right = tbl.rows[0].cells
        c_left.width = Cm(2.2)
        c_center.width = Cm(12.6)
        c_right.width = Cm(2.2)

        # Left logo
        p_left = c_left.paragraphs[0]
        p_left.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_left.paragraph_format.space_after = Pt(0)
        logo_u = self._assets_dir / "logo_unitbv.png"
        if logo_u.exists():
            try:
                p_left.add_run().add_picture(str(logo_u), width=Cm(1.8), height=Cm(1.8))
            except Exception:
                pass

        # Center text
        p_mid = c_center.paragraphs[0]
        p_mid.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_mid.paragraph_format.space_after = Pt(2)
        p_mid.paragraph_format.line_spacing = 1.05

        r1 = p_mid.add_run("UNIVERSITATEA TRANSILVANIA DIN BRAȘOV\n")
        self._apply_font(r1, size_pt=12.0, color_rgb=self.COLOR_NAVY, bold=True)
        r2 = p_mid.add_run("FACULTATEA DE INGINERIE ELECTRICĂ ȘI ȘTIINȚA CALCULATOARELOR\n")
        self._apply_font(r2, size_pt=9.5, color_rgb=self.COLOR_SLATE, bold=True)
        r3 = p_mid.add_run("Departamentul de Automatică și Tehnologia Informației")
        self._apply_font(r3, size_pt=9.0, color_rgb=self.COLOR_MUTED, italic=True)

        # Right logo
        p_right = c_right.paragraphs[0]
        p_right.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_right.paragraph_format.space_after = Pt(0)
        logo_i = self._assets_dir / "logo_iesc.png"
        if logo_i.exists():
            try:
                p_right.add_run().add_picture(str(logo_i), width=Cm(1.8), height=Cm(1.8))
            except Exception:
                pass

        # Divider line
        p_div = doc.add_paragraph()
        p_div.paragraph_format.space_before = Pt(6)
        p_div.paragraph_format.space_after = Pt(14)
        pPr = p_div._p.get_or_add_pPr()
        pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="12" w:space="1" w:color="{self.HEX_NAVY}"/></w:pBdr>')
        pPr.append(pBdr)

    # -------------------------------------------------------------------------
    # 1. CONVENȚIE-CADRU DE PRACTICĂ
    # -------------------------------------------------------------------------
    def generate_convention_docx(self, data: Optional[Dict[str, Any]] = None) -> bytes:
        """Generates the official UNITBV Convenție-cadru de practică DOCX file."""
        data = data or {}
        student_name = data.get("student_name", "Ciobănoiu Rareș-Alexandru")
        study_program = data.get("study_program", "Tehnologia Informației")
        study_year = data.get("study_year", "Anul II")
        group = data.get("group", "4LF341")
        academic_year = data.get("academic_year", settings.CURRENT_ACADEMIC_YEAR)
        hours_total = data.get("hours_total", 90)
        ects_credits = data.get("ects_credits", 4)
        company_name = data.get("company_name", "[Denumire Partener de Practică]")
        company_cui = data.get("company_cui", "[CUI/CIF]")
        company_rep = data.get("company_rep", "[Reprezentant Legal Partener]")
        company_address = data.get("company_address", "[Sediul social / Adresa companiei]")
        tutor_name = data.get("tutor_name", "[Nume și Prenume Tutore Practică]")
        tutor_role = data.get("tutor_role", "Senior Software Engineer / Mentor")
        supervisor_name = data.get("supervisor_name", "Șef Lucr. Dr. Ing. Dragoș Bratu")
        practice_period = data.get("practice_period", "15 Iulie 2026 – 28 August 2026")
        project_topic = data.get("project_topic", "Dezvoltarea unei Platforme Asistive AI pentru Management Academic și Practică")

        doc = docx.Document()
        self._configure_document_page(doc, "Convenție-cadru de Practică")
        self._add_institutional_header(doc)

        # Title
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_title.paragraph_format.space_before = Pt(4)
        p_title.paragraph_format.space_after = Pt(2)
        r_t1 = p_title.add_run("CONVENȚIE-CADRU\n")
        self._apply_font(r_t1, size_pt=14.0, color_rgb=self.COLOR_NAVY, bold=True)
        r_t2 = p_title.add_run("privind efectuarea stagiului de practică în cadrul programelor de studii universitare de licență\n")
        self._apply_font(r_t2, size_pt=11.0, color_rgb=self.COLOR_SLATE, bold=True)
        r_t3 = p_title.add_run(f"Anul Universitar {academic_year}")
        self._apply_font(r_t3, size_pt=10.0, color_rgb=self.COLOR_MUTED, italic=True)

        p_preamble = doc.add_paragraph()
        p_preamble.paragraph_format.space_before = Pt(10)
        p_preamble.paragraph_format.space_after = Pt(8)
        p_preamble.paragraph_format.line_spacing = 1.15
        p_preamble.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r_pre = p_preamble.add_run(
            "Prezenta convenție-cadru se încheie în conformitate cu prevederile Legii nr. 258/2007 privind practica elevilor și studenților, "
            "precum și cu Regulamentul-cadru al Universității Transilvania din Brașov privind organizarea și desfășurarea stagiilor de practică, "
            "între următoarele părți:"
        )
        self._apply_font(r_pre, size_pt=10.0, color_rgb=self.COLOR_DARK)

        # Parties Table
        parties_table = doc.add_table(rows=3, cols=2)
        parties_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        self._set_table_borders(parties_table, self.HEX_BORDER)

        parties_info = [
            ("1. Instituția de Învățământ Superior Organizatoare:", (
                "UNIVERSITATEA TRANSILVANIA DIN BRAȘOV (UNITBV)\n"
                "Facultatea de Inginerie Electrică și Știința Calculatoarelor (FIESC)\n"
                "Sediul: Brașov, Bd. Eroilor nr. 29 | Decanat: Str. Politehnicii nr. 1\n"
                f"Cadru didactic îndrumător: {supervisor_name}"
            )),
            ("2. Partenerul de Practică (Compania):", (
                f"Denumire: {company_name}\n"
                f"Sediul / Adresa: {company_address} | CUI / CIF: {company_cui}\n"
                f"Reprezentant legal: {company_rep}\n"
                f"Tutore de practică desemnat: {tutor_name} ({tutor_role})"
            )),
            ("3. Studentul Practicant:", (
                f"Nume și Prenume: {student_name}\n"
                f"Programul de studii: {study_program} | {study_year}, Grupa: {group}\n"
                f"Facultatea: FIESC, Universitatea Transilvania din Brașov"
            )),
        ]

        for idx, (role, desc) in enumerate(parties_info):
            row = parties_table.rows[idx]
            self._set_row_cant_split(row)
            c0, c1 = row.cells[0], row.cells[1]
            c0.width = Cm(5.0)
            c1.width = Cm(12.0)
            self._set_cell_background(c0, self.HEX_LIGHT_BG)
            self._set_cell_background(c1, "FFFFFF")
            self._set_cell_margins(c0, top=80, bottom=80, left=100, right=100)
            self._set_cell_margins(c1, top=80, bottom=80, left=100, right=100)

            p0 = c0.paragraphs[0]
            p0.paragraph_format.space_after = Pt(0)
            r0 = p0.add_run(role)
            self._apply_font(r0, size_pt=9.5, color_rgb=self.COLOR_NAVY, bold=True)

            p1 = c1.paragraphs[0]
            p1.paragraph_format.space_after = Pt(0)
            p1.paragraph_format.line_spacing = 1.15
            r1 = p1.add_run(desc)
            self._apply_font(r1, size_pt=9.5, color_rgb=self.COLOR_DARK)

        # Articles
        articles = [
            ("Art. 1. Obiectul Convenției", (
                f"Prezenta convenție stabilește cadrul general de colaborare între Universitate, Partenerul de practică și Student "
                f"în vederea desfășurării stagiului de practică tehnologică/profesională prevăzut în planul de învățământ, "
                f"având durata normată de {hours_total} de ore ({ects_credits} credite ECTS transferabile)."
            )),
            ("Art. 2. Statutul Practicantului", (
                "Pe toată durata stagiului de practică, practicantul își păstrează calitatea de student al Universității Transilvania din Brașov. "
                "Stagiul de practică nu echivalează cu încheierea unui contract individual de muncă, cu excepția situației în care părțile convin expres acest lucru."
            )),
            ("Art. 3. Perioada și Locul de Desfășurare", (
                f"Stagiul de practică se va desfășura în perioada {practice_period}, la sediul partenerului de practică sau în regim hibrid/remote, "
                f"sub directa îndrumare a tutorelui desemnat."
            )),
            ("Art. 4. Drepturile și Obligațiile Partenerului de Practică", (
                "a) Să desemneze un tutore de practică cu experiență tehnică adecvată pentru îndrumarea studentului;\n"
                "b) Să asigure instructajul de protecția muncii (SSM) și prevenirea incendiilor (PSI) înaintea începerii activității;\n"
                "c) Să pună la dispoziția studentului echipamentele, mediul tehnic și licențele necesare realizării proiectului;\n"
                "d) Să vizeze săptămânal caietul de practică și să completeze fișa de evaluare acordând o notă de la 1 la 10 la finalul stagiului."
            )),
            ("Art. 5. Drepturile și Obligațiile Studentului Practicant", (
                "a) Să respecte programul de lucru convenit și normele de securitate și disciplină ale companiei;\n"
                "b) Să păstreze confidențialitatea datelor, surselor și informațiilor proprietate ale partenerului de practică;\n"
                "c) Să completeze cu regularitate Caietul de practică cu descrierea activităților și modulelor dezvoltate;\n"
                "d) Să elaboreze proiectul tehnic de practică și să susțină colocviul de practică în sesiunea programată."
            )),
            ("Art. 6. Drepturile și Obligațiile Universității", (
                "a) Să asigure coordonarea științifică prin cadrul didactic îndrumător;\n"
                "b) Să monitorizeze concordanța activității practice cu programa analitică a domeniului de licență;\n"
                "c) Să organizeze colocviul de practică în fața unei comisii universitare de specialitate."
            )),
            ("Art. 7. Evaluarea și Notarea Stagiului", (
                "Evaluarea finală se realizează prin susținerea Colocviului de practică. "
                "Nota finală se calculează ponderat: 50% nota propusă de tutore pe fișa de evaluare a companiei + 50% nota acordată de comisia de colocviu. "
                "Condiția obligatorie de promovare este obținerea a minimum nota 5 (cinci) la ambele componente."
            )),
        ]

        for art_title, art_text in articles:
            p_art = doc.add_paragraph()
            p_art.paragraph_format.space_before = Pt(8)
            p_art.paragraph_format.space_after = Pt(2)
            p_art.paragraph_format.keep_with_next = True
            r_at = p_art.add_run(art_title)
            self._apply_font(r_at, size_pt=10.5, color_rgb=self.COLOR_NAVY, bold=True)

            p_txt = doc.add_paragraph()
            p_txt.paragraph_format.space_after = Pt(4)
            p_txt.paragraph_format.line_spacing = 1.15
            p_txt.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            r_tx = p_txt.add_run(art_text)
            self._apply_font(r_tx, size_pt=9.5, color_rgb=self.COLOR_DARK)

        # Anexa Pedagogică
        p_anx = doc.add_paragraph()
        p_anx.paragraph_format.space_before = Pt(12)
        p_anx.paragraph_format.space_after = Pt(4)
        p_anx.paragraph_format.keep_with_next = True
        r_ax = p_anx.add_run("ANEXĂ PEDAGOGICĂ: PROIECTUL ȘI TEMA DE PRACTICĂ")
        self._apply_font(r_ax, size_pt=11.0, color_rgb=self.COLOR_TEAL, bold=True)

        anx_table = doc.add_table(rows=3, cols=2)
        anx_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        self._set_table_borders(anx_table, self.HEX_BORDER)

        anx_rows = [
            ("Tema / Titlul Proiectului:", project_topic),
            ("Obiective Pedagogice și Competențe Vizate:", (
                "1. Proiectarea și implementarea de arhitecturi software modulare decuplate;\n"
                "2. Integrarea API-urilor externe asincrone (LLM, Calendar, IMAP/SMTP, Vector Store);\n"
                "3. Aplicarea standardelor de securitate (HMAC, Token Masking, Human-in-the-Loop);\n"
                "4. Containerizare cu Docker & Docker Compose și testare automată (pytest)."
            )),
            ("Tehnologii și Instrumente de Lucru:", "Python 3.14, FastAPI, PostgreSQL, Qdrant Vector DB, Ollama, Docker, Telegram Bot API"),
        ]

        for idx, (col0, col1) in enumerate(anx_rows):
            row = anx_table.rows[idx]
            self._set_row_cant_split(row)
            c0, c1 = row.cells[0], row.cells[1]
            c0.width = Cm(5.0)
            c1.width = Cm(12.0)
            self._set_cell_background(c0, self.HEX_LIGHT_BG)
            self._set_cell_margins(c0, top=70, bottom=70, left=100, right=100)
            self._set_cell_margins(c1, top=70, bottom=70, left=100, right=100)

            p0 = c0.paragraphs[0]
            p0.paragraph_format.space_after = Pt(0)
            r0 = p0.add_run(col0)
            self._apply_font(r0, size_pt=9.0, color_rgb=self.COLOR_NAVY, bold=True)

            p1 = c1.paragraphs[0]
            p1.paragraph_format.space_after = Pt(0)
            p1.paragraph_format.line_spacing = 1.15
            r1 = p1.add_run(col1)
            self._apply_font(r1, size_pt=9.0, color_rgb=self.COLOR_DARK)

        # Signatures
        p_sig_title = doc.add_paragraph()
        p_sig_title.paragraph_format.space_before = Pt(16)
        p_sig_title.paragraph_format.space_after = Pt(6)
        p_sig_title.paragraph_format.keep_with_next = True
        r_st = p_sig_title.add_run("Semnăturile Părților Contractante (3 exemplare originale):")
        self._apply_font(r_st, size_pt=10.0, color_rgb=self.COLOR_NAVY, bold=True)

        sig_table = doc.add_table(rows=2, cols=3)
        sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        self._set_table_borders(sig_table, self.HEX_BORDER)

        sig_headers = [
            "Universitatea Transilvania din Brașov\nDecan / Cadru Didactic Îndrumător",
            f"Partenerul de Practică\n{company_name}\nReprezentant / Tutore",
            f"Studentul Practicant\n\n{student_name}",
        ]
        for c_idx, text in enumerate(sig_headers):
            cell = sig_table.rows[0].cells[c_idx]
            cell.width = Cm(5.6)
            self._set_cell_background(cell, self.HEX_LIGHT_BG)
            self._set_cell_margins(cell, top=80, bottom=80, left=60, right=60)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            self._apply_font(r, size_pt=8.5, color_rgb=self.COLOR_NAVY, bold=True)

        for c_idx in range(3):
            cell = sig_table.rows[1].cells[c_idx]
            cell.width = Cm(5.6)
            self._set_cell_margins(cell, top=200, bottom=100, left=60, right=60)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run("\n\nSemnătura și Ștampila:\n_______________________")
            self._apply_font(r, size_pt=8.5, color_rgb=self.COLOR_MUTED)

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    # -------------------------------------------------------------------------
    # 2. CAIET DE PRACTICĂ / JURNAL DE ACTIVITATE
    # -------------------------------------------------------------------------
    def generate_logbook_docx(self, data: Optional[Dict[str, Any]] = None) -> bytes:
        """Generates the official UNITBV Caiet de practică (Weekly Activity Logbook & Tutor Evaluation)."""
        data = data or {}
        student_name = data.get("student_name", "Ciobănoiu Rareș-Alexandru")
        study_program = data.get("study_program", "Tehnologia Informației")
        study_year = data.get("study_year", "Anul II")
        group = data.get("group", "4LF341")
        academic_year = data.get("academic_year", settings.CURRENT_ACADEMIC_YEAR)
        hours_total = data.get("hours_total", 90)
        company_name = data.get("company_name", "[Denumire Partener de Practică]")
        tutor_name = data.get("tutor_name", "[Nume Tutore Practică]")
        tutor_role = data.get("tutor_role", "Senior Software Engineer / Team Lead")
        supervisor_name = data.get("supervisor_name", "Șef Lucr. Dr. Ing. Dragoș Bratu")
        project_topic = data.get("project_topic", "Dezvoltarea unei Platforme Asistive AI pentru Management Academic și Practică")

        doc = docx.Document()
        self._configure_document_page(doc, "Caiet de Practică Studențească")
        self._add_institutional_header(doc)

        # Title
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_title.paragraph_format.space_before = Pt(4)
        p_title.paragraph_format.space_after = Pt(2)
        r_t1 = p_title.add_run("CAIET DE PRACTICĂ STUDENȚEASCĂ\n")
        self._apply_font(r_t1, size_pt=14.0, color_rgb=self.COLOR_NAVY, bold=True)
        r_t2 = p_title.add_run("Jurnal de Activitate Săptămânală și Fișă de Evaluare a Competențelor\n")
        self._apply_font(r_t2, size_pt=11.0, color_rgb=self.COLOR_SLATE, bold=True)
        r_t3 = p_title.add_run(f"Anul Universitar {academic_year} • Volum normat: {hours_total} ore")
        self._apply_font(r_t3, size_pt=10.0, color_rgb=self.COLOR_MUTED, italic=True)

        # Student Identification Card
        id_tbl = doc.add_table(rows=4, cols=2)
        id_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        self._set_table_borders(id_tbl, self.HEX_BORDER)

        id_rows = [
            ("Student Practicant:", f"{student_name} | Program: {study_program} ({study_year}, Grupa {group})"),
            ("Instituție de Învățământ:", f"Universitatea Transilvania din Brașov • FIESC\nCadru didactic îndrumător: {supervisor_name}"),
            ("Partener de Practică:", f"{company_name}\nTutore desemnat: {tutor_name} ({tutor_role})"),
            ("Proiect / Tematică Stagiului:", project_topic),
        ]

        for idx, (col0, col1) in enumerate(id_rows):
            row = id_tbl.rows[idx]
            self._set_row_cant_split(row)
            c0, c1 = row.cells[0], row.cells[1]
            c0.width = Cm(5.0)
            c1.width = Cm(12.0)
            self._set_cell_background(c0, self.HEX_LIGHT_BG)
            self._set_cell_margins(c0, top=70, bottom=70, left=100, right=100)
            self._set_cell_margins(c1, top=70, bottom=70, left=100, right=100)

            p0 = c0.paragraphs[0]
            p0.paragraph_format.space_after = Pt(0)
            r0 = p0.add_run(col0)
            self._apply_font(r0, size_pt=9.0, color_rgb=self.COLOR_NAVY, bold=True)

            p1 = c1.paragraphs[0]
            p1.paragraph_format.space_after = Pt(0)
            p1.paragraph_format.line_spacing = 1.1
            r1 = p1.add_run(col1)
            self._apply_font(r1, size_pt=9.0, color_rgb=self.COLOR_DARK, bold=(idx == 0))

        # Weekly Activity Logs (3 weeks x 30 hours = 90 hours)
        weeks_data = [
            {
                "week": "SĂPTĂMÂNA 1 (Ore efectuate: 30 ore)",
                "days": [
                    ("Ziua 1 (6h)", "Instructaj SSM/PSI. Prezentarea companiei, a echipei și a cerințelor de arhitectură ale sistemului."),
                    ("Ziua 2 (6h)", "Configurarea mediului de dezvoltare: Docker, Poetry/pip, repo Git, conexiune PostgreSQL și Qdrant."),
                    ("Ziua 3 (6h)", "Proiectarea schemelor de date relaționale (tabele audit, memorie conversațională, pending drafts) și migrații Alembic."),
                    ("Ziua 4 (6h)", "Implementarea modulului de bază FastAPI, definirea endpoint-urilor de sănătate (/health/liveness, /health/readiness)."),
                    ("Ziua 5 (6h)", "Integrarea mecanismului de rate-limiting (SlidingWindow) și a verificării criptografice HMAC pentru securitatea webhook-ului."),
                ],
                "competences": "Setup mediu de lucru, arhitectură backend FastAPI, securizare API cu HMAC, proiectare baze de date.",
            },
            {
                "week": "SĂPTĂMÂNA 2 (Ore efectuate: 30 ore)",
                "days": [
                    ("Ziua 6 (6h)", "Implementarea pipeline-ului RAG de practică: parsare ghiduri UNITBV, chunking glisant și încărcare vectori în Qdrant."),
                    ("Ziua 7 (6h)", "Integrarea providerilor LLM hibrizi (OpenAI Cloud + Ollama Local on-premise) cu fallback determinist."),
                    ("Ziua 8 (6h)", "Conexiune IMAP/SMTP duală (cont personal + instituțional @student.unitbv.ro) și clasificare automată e-mailuri."),
                    ("Ziua 9 (6h)", "Implementarea fluxului Human-in-the-Loop: generare draft răspuns universitar condiționat de aprobare pe Telegram."),
                    ("Ziua 10 (6h)", "Integrarea Google Calendar API v3: autentificare Service Account, parsare limbaj natural și detecție conflicte orare."),
                ],
                "competences": "Căutare semantică vectorială (RAG), orchestrare agenți LLM, protocoale IMAP/SMTP, Google Calendar API.",
            },
            {
                "week": "SĂPTĂMÂNA 3 (Ore efectuate: 30 ore)",
                "days": [
                    ("Ziua 11 (6h)", "Implementarea sistemului de memorie conversațională multi-turn și extragerea automată a sarcinilor (Action Items)."),
                    ("Ziua 12 (6h)", "Dezvoltarea modulului NewsAgent cu parsare fluxuri RSS tehnologice și deduplicare criptografică SHA-256."),
                    ("Ziua 13 (6h)", "Generare automată de documente oficiale de practică în format DOCX (Convenție și Caiet de activitate) și livrare directă pe Telegram."),
                    ("Ziua 14 (6h)", "Scrierea suitei de teste automate (90+ teste unitare și de integrare pytest) și verificarea scenariilor E2E."),
                    ("Ziua 15 (6h)", "Finalizarea raportului tehnic de practică, revizuirea dosarului cu tutorele și pregătirea susținerii colocviului."),
                ],
                "competences": "Memorie conversațională, generare automată de documente, testare automată completă, finalizare dosar practică.",
            },
        ]

        for w in weeks_data:
            p_wh = doc.add_paragraph()
            p_wh.paragraph_format.space_before = Pt(12)
            p_wh.paragraph_format.space_after = Pt(4)
            p_wh.paragraph_format.keep_with_next = True
            r_wh = p_wh.add_run(w["week"])
            self._apply_font(r_wh, size_pt=11.0, color_rgb=self.COLOR_NAVY, bold=True)

            tbl_w = doc.add_table(rows=len(w["days"]) + 2, cols=2)
            tbl_w.alignment = WD_TABLE_ALIGNMENT.CENTER
            self._set_table_borders(tbl_w, self.HEX_BORDER)

            for d_idx, (day_label, day_desc) in enumerate(w["days"]):
                row = tbl_w.rows[d_idx]
                self._set_row_cant_split(row)
                c0, c1 = row.cells[0], row.cells[1]
                c0.width = Cm(3.5)
                c1.width = Cm(13.5)
                self._set_cell_background(c0, self.HEX_LIGHT_BG)
                self._set_cell_margins(c0, top=60, bottom=60, left=80, right=80)
                self._set_cell_margins(c1, top=60, bottom=60, left=80, right=80)

                p0 = c0.paragraphs[0]
                p0.paragraph_format.space_after = Pt(0)
                r0 = p0.add_run(day_label)
                self._apply_font(r0, size_pt=9.0, color_rgb=self.COLOR_SLATE, bold=True)

                p1 = c1.paragraphs[0]
                p1.paragraph_format.space_after = Pt(0)
                p1.paragraph_format.line_spacing = 1.15
                r1 = p1.add_run(day_desc)
                self._apply_font(r1, size_pt=9.0, color_rgb=self.COLOR_DARK)

            # Competences row
            row_comp = tbl_w.rows[len(w["days"])]
            self._set_row_cant_split(row_comp)
            c0, c1 = row_comp.cells[0], row_comp.cells[1]
            c0.width = Cm(3.5)
            c1.width = Cm(13.5)
            self._set_cell_background(c0, "EEF2F6")
            self._set_cell_background(c1, self.HEX_LIGHT_BG)
            self._set_cell_margins(c0, top=60, bottom=60, left=80, right=80)
            self._set_cell_margins(c1, top=60, bottom=60, left=80, right=80)

            p0 = c0.paragraphs[0]
            p0.paragraph_format.space_after = Pt(0)
            r0 = p0.add_run("Competențe:")
            self._apply_font(r0, size_pt=9.0, color_rgb=self.COLOR_NAVY, bold=True)

            p1 = c1.paragraphs[0]
            p1.paragraph_format.space_after = Pt(0)
            r1 = p1.add_run(w["competences"])
            self._apply_font(r1, size_pt=9.0, color_rgb=self.COLOR_DARK, italic=True)

            # Tutor sign row
            row_sign = tbl_w.rows[len(w["days"]) + 1]
            self._set_row_cant_split(row_sign)
            c0, c1 = row_sign.cells[0], row_sign.cells[1]
            c0.width = Cm(3.5)
            c1.width = Cm(13.5)
            self._set_cell_margins(c0, top=60, bottom=60, left=80, right=80)
            self._set_cell_margins(c1, top=60, bottom=60, left=80, right=80)

            p0 = c0.paragraphs[0]
            p0.paragraph_format.space_after = Pt(0)
            r0 = p0.add_run("Viza Tutorelui:")
            self._apply_font(r0, size_pt=9.0, color_rgb=self.COLOR_MUTED, bold=True)

            p1 = c1.paragraphs[0]
            p1.paragraph_format.space_after = Pt(0)
            r1 = p1.add_run("Activitate realizată integral (30 ore). Semnătură tutore: _______________________")
            self._apply_font(r1, size_pt=9.0, color_rgb=self.COLOR_MUTED)

        # Final Evaluation Sheet by Company Tutor
        p_ev = doc.add_paragraph()
        p_ev.paragraph_format.space_before = Pt(16)
        p_ev.paragraph_format.space_after = Pt(4)
        p_ev.paragraph_format.keep_with_next = True
        r_ev = p_ev.add_run("FIȘĂ DE EVALUARE ȘI APRECIERE SINTETICĂ A TUTORELUI DE PRACTICĂ")
        self._apply_font(r_ev, size_pt=11.0, color_rgb=self.COLOR_NAVY, bold=True)

        ev_tbl = doc.add_table(rows=6, cols=3)
        ev_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        self._set_table_borders(ev_tbl, self.HEX_BORDER)

        ev_headers = ["Nr.", "Criteriu de Evaluare a Studentului Practicant", "Notă Acordată (1 - 10)"]
        for c_idx, h_text in enumerate(ev_headers):
            cell = ev_tbl.rows[0].cells[c_idx]
            self._set_cell_background(cell, self.HEX_NAVY)
            self._set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(h_text)
            self._apply_font(r, size_pt=9.0, color_rgb=self.COLOR_WHITE, bold=True)

        ev_tbl.rows[0].cells[0].width = Cm(1.2)
        ev_tbl.rows[0].cells[1].width = Cm(11.8)
        ev_tbl.rows[0].cells[2].width = Cm(4.0)

        criteria = [
            ("1", "Disciplina, respectarea programului și a normelor de securitate a muncii", "10 (zece)"),
            ("2", "Asimilarea rapidă a cunoștințelor tehnologice și a uneltelor de lucru", "10 (zece)"),
            ("3", "Calitatea arhitecturii software, a codului sursă și a testelor implementate", "10 (zece)"),
            ("4", "Autonomie, capacitate de documentare tehnică și rezolvare a problemelor", "10 (zece)"),
            ("5", "Calitatea raportului tehnic și acuratețea jurnalului de activitate", "10 (zece)"),
        ]

        for idx, (nr, crit, nota) in enumerate(criteria, start=1):
            row = ev_tbl.rows[idx]
            self._set_row_cant_split(row)
            c0, c1, c2 = row.cells[0], row.cells[1], row.cells[2]
            c0.width = Cm(1.2)
            c1.width = Cm(11.8)
            c2.width = Cm(4.0)
            bg = self.HEX_LIGHT_BG if idx % 2 == 1 else "FFFFFF"
            self._set_cell_background(c0, bg)
            self._set_cell_background(c1, bg)
            self._set_cell_background(c2, bg)
            self._set_cell_margins(c0, top=70, bottom=70, left=60, right=60)
            self._set_cell_margins(c1, top=70, bottom=70, left=80, right=80)
            self._set_cell_margins(c2, top=70, bottom=70, left=60, right=60)

            p0 = c0.paragraphs[0]
            p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p0.paragraph_format.space_after = Pt(0)
            r0 = p0.add_run(nr)
            self._apply_font(r0, size_pt=9.0, color_rgb=self.COLOR_DARK)

            p1 = c1.paragraphs[0]
            p1.paragraph_format.space_after = Pt(0)
            r1 = p1.add_run(crit)
            self._apply_font(r1, size_pt=9.0, color_rgb=self.COLOR_DARK)

            p2 = c2.paragraphs[0]
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p2.paragraph_format.space_after = Pt(0)
            r2 = p2.add_run(nota)
            self._apply_font(r2, size_pt=9.0, color_rgb=self.COLOR_NAVY, bold=True)

        # Final grade and conclusion
        p_concl = doc.add_paragraph()
        p_concl.paragraph_format.space_before = Pt(12)
        p_concl.paragraph_format.space_after = Pt(4)
        p_concl.paragraph_format.line_spacing = 1.15
        r_cg = p_concl.add_run("NOTĂ FINALĂ PROPUSĂ DE TUTORELE COMPANIEI PENTRU COLOCVIU:  10 (zece)\n")
        self._apply_font(r_cg, size_pt=10.5, color_rgb=self.COLOR_NAVY, bold=True)
        r_ap = p_concl.add_run(
            "Aprecieri calitative ale tutorelui: Studentul a demonstrat un nivel excelent de pregătire teoretică și aplicativă, "
            "implementând cu rigurozitate cerințele platformei backend, respectând cele mai bune practici de securitate și decuplare modulară."
        )
        self._apply_font(r_ap, size_pt=9.5, color_rgb=self.COLOR_DARK, italic=True)

        # Signatures table
        p_fsig = doc.add_paragraph()
        p_fsig.paragraph_format.space_before = Pt(14)
        p_fsig.paragraph_format.space_after = Pt(6)
        p_fsig.paragraph_format.keep_with_next = True
        r_fst = p_fsig.add_run("Validare și Semnături Finale:")
        self._apply_font(r_fst, size_pt=10.0, color_rgb=self.COLOR_NAVY, bold=True)

        sig_ev_tbl = doc.add_table(rows=2, cols=2)
        sig_ev_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        self._set_table_borders(sig_ev_tbl, self.HEX_BORDER)

        sig_ev_tbl.rows[0].cells[0].width = Cm(8.5)
        sig_ev_tbl.rows[0].cells[1].width = Cm(8.5)
        sig_ev_tbl.rows[1].cells[0].width = Cm(8.5)
        sig_ev_tbl.rows[1].cells[1].width = Cm(8.5)

        c0_title = sig_ev_tbl.rows[0].cells[0]
        self._set_cell_background(c0_title, self.HEX_LIGHT_BG)
        p = c0_title.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(f"Tutore de Practică Companie\n{tutor_name} ({company_name})")
        self._apply_font(r, size_pt=9.0, color_rgb=self.COLOR_NAVY, bold=True)

        c1_title = sig_ev_tbl.rows[0].cells[1]
        self._set_cell_background(c1_title, self.HEX_LIGHT_BG)
        p = c1_title.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(f"Student Practicant\n{student_name}")
        self._apply_font(r, size_pt=9.0, color_rgb=self.COLOR_NAVY, bold=True)

        c0_sig = sig_ev_tbl.rows[1].cells[0]
        self._set_cell_margins(c0_sig, top=160, bottom=80, left=60, right=60)
        p = c0_sig.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run("Semnătura și Ștampila Companiei:\n_______________________")
        self._apply_font(r, size_pt=8.5, color_rgb=self.COLOR_MUTED)

        c1_sig = sig_ev_tbl.rows[1].cells[1]
        self._set_cell_margins(c1_sig, top=160, bottom=80, left=60, right=60)
        p = c1_sig.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run("Semnătura Studentului:\n_______________________")
        self._apply_font(r, size_pt=8.5, color_rgb=self.COLOR_MUTED)

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    # -------------------------------------------------------------------------
    # 3. CONVENIENT PACKAGE GENERATION
    # -------------------------------------------------------------------------
    def generate_practice_package(self, data: Optional[Dict[str, Any]] = None) -> Dict[str, bytes]:
        """Generates both Convenție and Caiet de practică."""
        return {
            "Conventie_Cadru_Practica_UNITBV.docx": self.generate_convention_docx(data),
            "Caiet_de_Practica_UNITBV.docx": self.generate_logbook_docx(data),
        }
