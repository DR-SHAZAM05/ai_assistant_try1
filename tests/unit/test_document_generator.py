import io
import docx
import pytest
from src.app.services.document_generator_service import DocumentGeneratorService
from src.app.orchestrator.orchestrator import AIOrchestrator
from src.app.orchestrator.intent import IntentType


def test_generate_convention_docx():
    service = DocumentGeneratorService()
    custom_data = {
        "student_name": "Test Student UNITBV",
        "study_program": "Tehnologia Informației",
        "group": "4LF341",
        "company_name": "Bitdefender România",
        "tutor_name": "Ing. Popescu Andrei",
    }
    doc_bytes = service.generate_convention_docx(custom_data)

    assert isinstance(doc_bytes, bytes)
    assert len(doc_bytes) > 5000

    # Parse resulting document with python-docx
    doc = docx.Document(io.BytesIO(doc_bytes))
    full_text = "\n".join([p.text for p in doc.paragraphs])
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                full_text += "\n" + cell.text

    assert "CONVENȚIE-CADRU" in full_text
    assert "Test Student UNITBV" in full_text
    assert "Bitdefender România" in full_text
    assert "Ing. Popescu Andrei" in full_text
    assert "UNIVERSITATEA TRANSILVANIA DIN BRAȘOV" in full_text


def test_generate_logbook_docx():
    service = DocumentGeneratorService()
    custom_data = {
        "student_name": "Test Student UNITBV",
        "study_program": "Tehnologia Informației",
        "group": "4LF341",
        "company_name": "Siemens Brașov",
        "tutor_name": "Ing. Ionescu Radu",
    }
    doc_bytes = service.generate_logbook_docx(custom_data)

    assert isinstance(doc_bytes, bytes)
    assert len(doc_bytes) > 5000

    doc = docx.Document(io.BytesIO(doc_bytes))
    full_text = "\n".join([p.text for p in doc.paragraphs])
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                full_text += "\n" + cell.text

    assert "CAIET DE PRACTICĂ STUDENȚEASCĂ" in full_text
    assert "Test Student UNITBV" in full_text
    assert "Siemens Brașov" in full_text
    assert "SĂPTĂMÂNA 1" in full_text
    assert "SĂPTĂMÂNA 2" in full_text
    assert "SĂPTĂMÂNA 3" in full_text
    assert "FIȘĂ DE EVALUARE" in full_text


def test_generate_practice_package():
    service = DocumentGeneratorService()
    package = service.generate_practice_package()

    assert "Conventie_Cadru_Practica_UNITBV.docx" in package
    assert "Caiet_de_Practica_UNITBV.docx" in package
    assert len(package["Conventie_Cadru_Practica_UNITBV.docx"]) > 5000
    assert len(package["Caiet_de_Practica_UNITBV.docx"]) > 5000


@pytest.mark.asyncio
async def test_orchestrator_practice_document_intent():
    orchestrator = AIOrchestrator()

    # Test Intent Detection
    res1 = await orchestrator.detect_intent("Generează convenția de practică")
    assert res1.intent == IntentType.PRACTICE_DOCUMENT_REQUEST

    res2 = await orchestrator.detect_intent("Vreau caietul de practică în format docx")
    assert res2.intent == IntentType.PRACTICE_DOCUMENT_REQUEST

    res3 = await orchestrator.detect_intent("Descarcă documentele de practică")
    assert res3.intent == IntentType.PRACTICE_DOCUMENT_REQUEST

    # Test Execution produces attached documents
    exec_result = await orchestrator.process_request("Generează convenția de practică și caietul")
    assert exec_result["intent"] == IntentType.PRACTICE_DOCUMENT_REQUEST.value
    assert "documents" in exec_result
    assert len(exec_result["documents"]) == 2
    filenames = [d["filename"] for d in exec_result["documents"]]
    assert "Conventie_Cadru_Practica_UNITBV.docx" in filenames
    assert "Caiet_de_Practica_UNITBV.docx" in filenames
